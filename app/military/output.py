"""Readable event outputs; diagnostics stay in JSON, never Word."""
import re

LEVEL_NAMES = {"major": "重大军武事件", "important": "重要军武动态", "normal": "一般军武新闻",
               "low": "低价值内容", "factual": "相关事实记录（不评分）"}


def _safe(text):
    return re.sub(r"([\\`*_{}\[\]<>#])", r"\\\1", str(text)).replace("\n", " ").replace("\r", " ")


def markdown(events):
    lines = ["# 军武动态", "", "依据已有新闻标题和摘要自动整理；分数为事件编排参考，具体事实以原报道为准。", ""]
    if not events:
        return "\n".join(lines + ["本时间段没有识别到涉台军武事件。", ""])
    for level, name in LEVEL_NAMES.items():
        selected = [e for e in events if e["level"] == level]
        if not selected:
            continue
        lines.extend([f"## {name}", ""])
        for idx, event in enumerate(selected, 1):
            rep = event["representative"]
            score = event["importance"] if event["importance"] is not None else "不评分"
            lines.extend([f"### {idx}. {_safe(event['title'])}", "",
                          f"类别：{event['category_name']}  ", f"重要度：{score}  ",
                          f"时间：{event['first_seen']} 至 {event['last_updated']}  ",
                          f"来源：{_safe('、'.join(event['sources']))}  ",
                          f"核心实体：{_safe('、'.join(event['core_entities']))}  ",
                          f"代表报道：[{_safe(rep['title'])}](<{rep['url'].replace('>', '%3E')}>)", "",
                          "研判依据：" + "；".join(event["importance_reasons"]), "",
                          f"相关新闻：{event['article_count']} 篇", ""])
            for row in event["articles"]:
                lines.append(f"- [{_safe(row['source'])}：{_safe(row['title'])}](<{row['url'].replace('>', '%3E')}>)")
            lines.append("")
    return "\n".join(lines)


def render_word_events(doc, events, catch_up_urls=None):
    from docx.shared import Pt
    from ..word_digest import _add_hyperlink

    for idx, event in enumerate(events, 1):
        rep = event["representative"]
        catch_up = any(a["url"] in (catch_up_urls or set()) for a in event["articles"])
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{idx}. {'【补发】' if catch_up else ''}{event['title']}")
        run.bold = True
        run.font.size = Pt(12)
        score = event["importance"] if event["importance"] is not None else "不评分"
        p = doc.add_paragraph(f"类别：{event['category_name']} ｜ 重要度：{score} ｜ 代表来源：{rep['source']} ｜ 相关新闻：{event['article_count']} 篇")
        p.paragraph_format.keep_with_next = True
        p = doc.add_paragraph("为何重要：" + "；".join(event["importance_reasons"]))
        p.paragraph_format.keep_with_next = True
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        _add_hyperlink(p, rep["url"], "代表报道")
        doc.add_paragraph()
