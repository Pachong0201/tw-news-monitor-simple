from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest

from app.article_identity import (
    article_identity_key,
    deduplicate_articles_by_identity,
    prefer_richer_article,
)
from app.collectors.base import BaseCollector
from app.collectors.military import (
    MNAMilitaryCollector,
    NownewsMilitaryCollector,
    _parse_ltn_date,
)
from app.database import Database, SCHEMA_VERSION
from app.election2026.config import load_election_config, load_entities
from app.freshness import filter_fresh_articles
from app.importance import classify_articles
from app.content_filter import (
    apply_content_filter,
    filter_mode,
)
from app.main import (
    _classify_delivery_articles,
    _get_source_continuity,
    collect_all,
    deduplicate_articles_by_url,
)
from app.news_pipeline import run_delivery_core
from app.source_health import SourceHealthStore, SourceOutcome
from app.military import (
    _military_source_ids_cached,
    clear_military_source_cache,
    military_source_ids,
)
from app.models import Article
from app.settings import get_settings
from app.topic_backfill import _safe_copy_for_dry_run, backfill_topics
from app.word_digest import build_word_digest
from docx import Document


TAIPEI = ZoneInfo("Asia/Taipei")


def article(
    url: str,
    title: str = "测试新闻",
    *,
    source_id: str = "test",
    category: str = "politics",
    published_at: datetime | None = None,
    precision: str = "exact",
) -> Article:
    now = datetime(2026, 9, 8, 12, 0, tzinfo=TAIPEI)
    return Article(
        source_id=source_id,
        source_name="测试来源",
        category=category,
        title=title,
        url=url,
        published_at=published_at,
        fetched_at=now,
        position=1,
        published_at_precision=precision,
    )


def test_importance_prefers_level_then_score_and_uses_summary():
    rules = {
        "enabled": True,
        "thresholds": {"critical": 85, "important": 65, "normal": 0},
        "scoring": {"multi_rule_bonus": 0},
        "rules": [
            {
                "id": "critical_security",
                "track": "politics_security",
                "base_score": 86,
                "level_cap": "critical",
                "subjects": ["賴清德"],
                "actions": ["國防改革"],
                "scenes": [],
                "negative": [],
            },
            {
                "id": "important_policy",
                "track": "politics_security",
                "base_score": 95,
                "level_cap": "important",
                "subjects": ["國家安全"],
                "actions": ["政策"],
                "scenes": [],
                "negative": [],
            },
        ],
    }
    title = "賴清德今日出席活動"
    summary = "宣布新的國家安全與兩岸政策，並提出國防改革方案。"
    title_only = article("https://example.test/importance/title", title, published_at=None)
    title_only.summary = None
    title_result = classify_articles([title_only], rules)[0][1]
    assert title_result.level == "normal"

    full = article("https://example.test/importance/full", title)
    full.summary = summary
    result = classify_articles([full], rules)[0][1]
    assert result.level == "critical"
    assert result.score == 86

    reversed_rules = dict(rules)
    reversed_rules["rules"] = list(reversed(rules["rules"]))
    reversed_result = classify_articles([full], reversed_rules)[0][1]
    assert (reversed_result.level, reversed_result.score) == ("critical", 86)


def test_display_url_preserves_path_case_and_removes_tracking():
    raw = "https://example.com/News/ABC123?utm_source=test&fbclid=x#section"
    display = BaseCollector.normalize_url(raw)
    assert display == "https://example.com/News/ABC123"
    assert article_identity_key(raw) == article_identity_key(display)


def test_military_time_precision_and_rollover():
    now = datetime(2026, 9, 8, 0, 10, tzinfo=TAIPEI)
    parsed = _parse_ltn_date("更新時間 23:58", now)
    assert parsed == datetime(2026, 9, 7, 23, 58, tzinfo=TAIPEI)


def test_nownews_aware_and_naive_timestamps():
    collector = NownewsMilitaryCollector({
        "id": "nownews_military",
        "name": "NOWnews",
        "category": "military",
        "url": "https://www.nownews.com/cat/news-summary/military/",
    })
    html = """
    <ul id="ulNewsList">
      <li class="item"><a href="/news/1"><h3 class="title">aware</h3>
        <time datetime="2026-09-08T00:00:00+00:00">x</time></a></li>
      <li class="item"><a href="/news/2"><h3 class="title">naive</h3>
        <time datetime="2026-09-08T00:00:00">x</time></a></li>
    </ul>
    """
    rows, valid = collector._parse_page(html, datetime(2026, 9, 8, 12, tzinfo=TAIPEI))
    assert valid
    by_title = {row.title: row for row in rows}
    assert by_title["aware"].published_at == datetime(2026, 9, 8, 8, tzinfo=TAIPEI)
    assert by_title["naive"].published_at == datetime(2026, 9, 8, 0, tzinfo=TAIPEI)


def test_mna_date_only_is_not_exact_time():
    class Client:
        def get(self, url):
            import httpx
            request = httpx.Request("GET", url)
            return httpx.Response(
                200,
                request=request,
                text=(
                    '<a href="/news/detail/?UserKey=01234567-89ab-cdef-0123-456789abcdef">'
                    '<div class="title">date only</div>'
                    '<div class="summary">summary</div>'
                    '<span class="time">民國115年09月08日</span></a>'
                ),
            )

    collector = MNAMilitaryCollector({
        "id": "mna_military",
        "name": "MNA",
        "category": "military",
        "url": "https://mna.mnd.gov.tw/news/overview/",
    })
    collector._client = Client()
    rows = collector.collect()
    assert len(rows) == 1
    assert rows[0].published_at_precision == "date_only"
    result = filter_fresh_articles(
        rows,
        datetime(2026, 9, 8, 12, tzinfo=TAIPEI),
    )
    # date_only same-day is a delivery candidate, not an unknown-time item.
    assert result.date_only_today_articles == rows
    assert result.unknown_time_articles == []


def test_schema_identity_and_batch_topics(tmp_path):
    db = Database(tmp_path / "news.db")
    db.connect()
    try:
        assert db.schema_version() == SCHEMA_VERSION
        columns = {row[1] for row in db.conn.execute("PRAGMA table_info(articles)")}
        assert {"identity_key", "published_at_precision"} <= columns
        assert db.save_articles([article("https://example.test/News/1")])
        row = db.conn.execute(
            "SELECT identity_key, published_at_precision FROM articles"
        ).fetchone()
        assert row[0] == article_identity_key("https://example.test/News/1")
        assert row[1] == "exact"
        db.save_topics([
            {"url": "https://example.test/News/1", "topic": "military", "source_type": "commercial_military"},
            {"url": "https://example.test/News/1", "topic": "election_2026_local", "scope": "national"},
        ])
        assert db.count_topics() == 2
        db.save_topics([{"url": "https://example.test/News/1", "topic": "military", "source_type": "official_military"}])
        assert db.count_topics() == 2
    finally:
        db.close()


def test_topic_ingest_happens_before_freshness(tmp_path):
    class Collector:
        def __init__(self, cfg):
            self.cfg = cfg

        def collect(self):
            old = datetime(2026, 8, 1, 12, tzinfo=TAIPEI)
            return [
                article("https://example.test/old", "2026九合一提名", published_at=old),
                article("https://example.test/unknown", "2026九合一提名", published_at=None),
            ]

        def close(self):
            pass

    db = Database(tmp_path / "topics.db")
    db.connect()
    try:
        result = collect_all(
            [{"id": "s", "name": "测试", "type": "stub", "category": "politics", "url": "https://example.test"}],
            db,
            collector_map={"stub": Collector},
            election_config=load_election_config(),
            election_entities=load_entities(),
        )
        assert len(result.inserted_articles) == 2
        assert db.get_topic_urls(
            ["https://example.test/old", "https://example.test/unknown"],
            "election_2026_local",
        ) == {"https://example.test/old", "https://example.test/unknown"}
    finally:
        db.close()


def test_topic_backfill_is_idempotent_and_dry_run_is_read_only(tmp_path):
    path = tmp_path / "backfill.db"
    db = Database(path)
    db.connect()
    db.save_articles([
        article("https://example.test/e", "2026九合一提名"),
        article("https://example.test/m", "军武新闻", source_id="mna_military", category="military"),
    ])
    db.close()
    cfg = load_election_config()
    entities = load_entities()
    sources = {"mna_military": {"topic": "military", "military_source_type": "official_military"}}

    ro = Database(path, read_only=True)
    ro.connect()
    try:
        preview = backfill_topics(
            ro, sources=sources, election_config=cfg, election_entities=entities,
            dry_run=True, batch_size=1,
        )
        assert preview["written"] == 0
        assert ro.count_topics() == 0
    finally:
        ro.close()

    db = Database(path)
    db.connect()
    try:
        first = backfill_topics(
            db, sources=sources, election_config=cfg, election_entities=entities,
            batch_size=1,
        )
        assert first["inserted"] >= 2
        assert first["updated"] == 0
        assert db.count_topics() >= 2

        select_row = (
            "SELECT scope, region, regions_json, event_type, confidence, "
            "source_type, metadata_json, created_at FROM news_topics "
            "WHERE url = ? AND topic = ?"
        )
        before = {
            (url, topic): db.conn.execute(select_row, (url, topic)).fetchone()
            for url in ("https://example.test/e", "https://example.test/m")
            for topic in ("election_2026_local", "military")
            if db.conn.execute(
                "SELECT 1 FROM news_topics WHERE url=? AND topic=?",
                (url, topic),
            ).fetchone()
        }
        count = db.count_topics()

        second = backfill_topics(
            db, sources=sources, election_config=cfg, election_entities=entities,
            batch_size=1,
        )
        assert second["inserted"] == 0
        assert second["updated"] == 0
        assert second["skipped_existing"] >= len(before)
        assert db.count_topics() == count
        after = {
            (url, topic): db.conn.execute(select_row, (url, topic)).fetchone()
            for url in ("https://example.test/e", "https://example.test/m")
            for topic in ("election_2026_local", "military")
            if db.conn.execute(
                "SELECT 1 FROM news_topics WHERE url=? AND topic=?",
                (url, topic),
            ).fetchone()
        }
        assert after == before
    finally:
        db.close()


def test_settings_path_priority_and_log_level(monkeypatch, tmp_path):
    for name in (
        "NEWS_DB_PATH", "DATABASE_PATH", "SOURCES_CONFIG_PATH", "LOG_LEVEL",
    ):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setenv("DATABASE_PATH", "legacy.db")
    monkeypatch.setenv("LOG_LEVEL", "debug")
    first = get_settings(tmp_path, load_env=False)
    assert first.news_db_path == tmp_path / "legacy.db"
    assert first.log_level == "DEBUG"
    monkeypatch.setenv("NEWS_DB_PATH", "new.db")
    second = get_settings(tmp_path, load_env=False)
    assert second.news_db_path == tmp_path / "new.db"


def test_military_source_cache_uses_custom_path_once(tmp_path):
    path = tmp_path / "sources.yaml"
    path.write_text(
        "sources:\n  - id: custom_military\n    topic: military\n    enabled: true\n",
        encoding="utf-8",
    )
    clear_military_source_cache()
    assert military_source_ids(path) == frozenset({"custom_military"})
    assert military_source_ids(path) == frozenset({"custom_military"})
    info = _military_source_ids_cached.cache_info()
    assert info.hits >= 1


def _make_dated(url, day_of_sep_2026, source_id="mna_military", title="MNA date only"):
    return article(
        url,
        title=title,
        source_id=source_id,
        category="military",
        published_at=datetime(2026, 9, day_of_sep_2026, tzinfo=TAIPEI),
        precision="date_only",
    )


def test_date_only_today_delivery_and_yesterday_not(tmp_path):
    today = datetime(2026, 9, 9, 14, 0, tzinfo=TAIPEI)
    today_article = _make_dated("https://mna.test/9", 9)
    yesterday_article = _make_dated("https://mna.test/8", 8)

    fresh_result = filter_fresh_articles([today_article], today)
    assert fresh_result.date_only_today_articles == [today_article]
    assert fresh_result.unknown_time_articles == []

    stale_result = filter_fresh_articles([yesterday_article], today)
    assert stale_result.stale_articles == [yesterday_article]
    assert stale_result.date_only_today_articles == []

    # Source baseline > 0 means this is a normal ongoing MNA run.
    delivery = _classify_delivery_articles(
        [today_article, yesterday_article],
        {"mna_military": 5},
        today,
        catch_up_enabled=False,
        source_continuity={"mna_military": True},
    )
    assert delivery["date_only_eligible"] == [today_article]
    assert delivery["stale_articles"] == [yesterday_article]


def test_date_only_baseline_does_not_flood_on_first_enable():
    run = datetime(2026, 9, 9, 14, 0, tzinfo=TAIPEI)
    baseline_articles = [
        _make_dated(f"https://mna.test/{i}", 9, title=f"baseline-{i}")
        for i in range(30)
    ]
    delivery = _classify_delivery_articles(
        baseline_articles,
        {"mna_military": 0},
        run,
        catch_up_enabled=False,
    )
    assert delivery["date_only_eligible"] == []
    assert len(delivery["stale_articles"]) == 30
    assert len(delivery["date_only_today"]) == 30


def test_unknown_precision_remains_not_delivered():
    run = datetime(2026, 9, 9, 14, 0, tzinfo=TAIPEI)
    unknown = article(
        "https://mna.test/unknown", "unknown", published_at=None, precision="unknown"
    )
    delivery = _classify_delivery_articles(
        [unknown], {"mna_military": 5}, run, catch_up_enabled=False
    )
    assert delivery["unknown_articles"] == [unknown]
    assert delivery["date_only_eligible"] == []
    assert delivery["fresh_articles"] == []


@pytest.mark.parametrize("reverse", [False, True])
def test_nownews_url_dedup_keeps_exact_metadata_regardless_order(reverse):
    unknown = article(
        "https://www.nownews.com/news/1",
        "NOWnews banner",
        category="military",
        published_at=None,
        precision="unknown",
    )
    exact = article(
        "https://www.nownews.com/news/1",
        "NOWnews list",
        category="military",
        published_at=datetime(2026, 9, 9, 18, 30, tzinfo=TAIPEI),
        precision="exact",
    )
    pair = [exact, unknown] if reverse else [unknown, exact]
    unique, dups = deduplicate_articles_by_url(pair)
    assert len(unique) == 1
    assert len(dups) == 1
    assert unique[0].published_at_precision == "exact"
    assert unique[0].published_at == datetime(2026, 9, 9, 18, 30, tzinfo=TAIPEI)


def test_identity_dedup_keeps_exact_metadata_both_orders():
    # Same identity (UDN story id) with one richer copy.
    unknown = article(
        "https://udn.com/news/story/6656/9635664",
        "alias banner",
        published_at=None,
        precision="unknown",
    )
    exact = article(
        "https://udn.com/news/story/124948/9635664",
        "alias exact",
        published_at=datetime(2026, 9, 9, 12, 0, tzinfo=TAIPEI),
        precision="exact",
    )
    for pair in ([unknown, exact], [exact, unknown]):
        unique, _ = deduplicate_articles_by_identity(pair)
        assert unique[0].published_at_precision == "exact"
        assert unique[0].published_at is not None


def _collect_with_mode(db, config):
    class StubCollector:
        def __init__(self, source):
            self.cfg = source

        def collect(self):
            return [
                article(
                    "https://example.test/a",
                    "正常政治新聞",
                    category="politics",
                    published_at=datetime(2026, 9, 9, 13, 0, tzinfo=TAIPEI),
                ),
                article(
                    "https://example.test/b",
                    "大樂透開獎",
                    category="economy",
                    published_at=datetime(2026, 9, 9, 13, 0, tzinfo=TAIPEI),
                ),
            ]

        def close(self):
            pass

    source = {
        "id": "stub", "name": "測試", "type": "stub",
        "category": "politics", "url": "https://example.test",
    }
    return collect_all(
        [source], db, config,
        collector_map={"stub": StubCollector},
    )


def test_delivery_eligibility_persistence_modes(tmp_path):
    db = Database(tmp_path / "eligibility.db")
    db.connect()
    try:
        # exclude_from_delivery: both saved, only one eligible.
        exclude = _collect_with_mode(
            db, {"enabled": True, "mode": "exclude_from_delivery",
                 "categories": {"_default": ["大樂透"]}}
        )
        assert len(exclude.inserted_articles) == 2
        assert len(exclude.filtered_from_delivery) == 1
        eligible = db.get_articles_since(datetime(2000, 1, 1), eligible_only=True)
        assert [a.title for a in eligible] == ["正常政治新聞"]
        row = db.conn.execute(
            "SELECT delivery_eligible, filter_reason FROM articles WHERE url=?",
            ("https://example.test/b",),
        ).fetchone()
        assert row[0] == 0
        assert row[1] == "content_filter"

        # drop_before_save on a fresh DB: only one persisted.
        db2 = Database(tmp_path / "drop.db")
        db2.connect()
        try:
            drop = _collect_with_mode(
                db2, {"enabled": True, "mode": "drop_before_save",
                      "categories": {"_default": ["大樂透"]}}
            )
            assert len(drop.inserted_articles) == 1
            assert len(drop.filtered_before_save) == 1
            assert db2.count_articles() == 1
        finally:
            db2.close()

        # disabled: all persisted and all eligible.
        db3 = Database(tmp_path / "disabled.db")
        db3.connect()
        try:
            disabled = _collect_with_mode(
                db3, {"enabled": False, "mode": "exclude_from_delivery",
                      "categories": {"_default": ["大樂透"]}}
            )
            assert len(disabled.inserted_articles) == 2
            assert db3.count_articles() == 2
            assert len(db3.get_articles_since(datetime(2000, 1, 1), eligible_only=True)) == 2
        finally:
            db3.close()
    finally:
        db.close()


def test_old_rows_default_delivery_eligible(tmp_path):
    import sqlite3
    path = tmp_path / "old.db"
    conn = sqlite3.connect(path)
    conn.execute("""
        CREATE TABLE articles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_id TEXT NOT NULL,
            source_name TEXT NOT NULL,
            category TEXT NOT NULL,
            title TEXT NOT NULL,
            url TEXT NOT NULL UNIQUE,
            published_at TEXT,
            fetched_at TEXT NOT NULL,
            position INTEGER NOT NULL
        )
    """)
    conn.execute(
        "INSERT INTO articles (source_id, source_name, category, title, url, "
        "published_at, fetched_at, position) VALUES (?,?,?,?,?,?,?,?)",
        ("s1", "源", "politics", "舊聞", "https://example.test/old",
         "2026-01-01T00:00:00", "2026-01-01T00:00:00", 1),
    )
    conn.commit()
    conn.close()

    db = Database(path)
    db.connect()
    try:
        cols = {row[1] for row in db.conn.execute("PRAGMA table_info(articles)")}
        assert {"delivery_eligible", "filter_reason", "filter_version"} <= cols
        row = db.conn.execute(
            "SELECT delivery_eligible, filter_reason, filter_version FROM articles"
        ).fetchone()
        assert row == (1, None, None)
        assert len(db.get_articles_since(datetime(2000, 1, 1), eligible_only=True)) == 1
    finally:
        db.close()


def test_invalid_content_filter_mode_falls_back_to_exclude(caplog):
    cfg = {"enabled": True, "mode": "typo_value",
           "categories": {"_default": ["大樂透"]}}
    mode = filter_mode(cfg)
    assert mode == "exclude_from_delivery"
    arts = [
        article("https://example.test/lot", "大樂透開獎", category="economy"),
        article("https://example.test/normal", "正常", category="politics"),
    ]
    result = apply_content_filter(arts, cfg)
    assert len(result.kept) == 1
    assert result.filtered_from_delivery == [arts[0]]
    assert "exclude_from_delivery" in caplog.text


def test_word_official_military_counts_are_self_consistent(tmp_path):
    now = datetime(2026, 9, 9, 12, 0, tzinfo=TAIPEI)
    official = article(
        "https://mna.mnd.gov.tw/news/detail?UserKey=abc",
        "国防部军闻",
        source_id="mna_military",
        category="military",
        published_at=now,
        precision="date_only",
    )
    output = build_word_digest(
        [official], tmp_path, generated_at=now,
        military_topic_urls={official.url},
    )
    text = "\n".join(p.text for p in Document(output).paragraphs)
    assert "新闻总数：1条" in text
    assert "官方信源：1条" in text
    assert "新闻媒体：0条" in text
    assert "军武动态" in text


def _create_old_articles_db(path, rows):
    import sqlite3
    conn = sqlite3.connect(str(path))
    conn.execute("""
        CREATE TABLE articles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_id TEXT NOT NULL,
            source_name TEXT NOT NULL,
            category TEXT NOT NULL,
            title TEXT NOT NULL,
            url TEXT NOT NULL UNIQUE,
            published_at TEXT,
            fetched_at TEXT NOT NULL,
            position INTEGER NOT NULL,
            summary TEXT,
            summary_source TEXT,
            summary_attempted_at TEXT
        )
    """)
    for i, row in enumerate(rows, 1):
        conn.execute(
            "INSERT INTO articles (source_id, source_name, category, title, url, "
            "published_at, fetched_at, position) VALUES (?,?,?,?,?,?,?,?)",
            (row["source_id"], row["source_name"], row["category"], row["title"],
             row["url"], row["published_at"], row["fetched_at"], i),
        )
    conn.commit()
    conn.close()


def test_topic_backfill_dry_run_old_schema_copy(tmp_path):
    path = tmp_path / "old-backfill.db"
    _create_old_articles_db(path, [
        {"source_id": "s", "source_name": "S", "category": "politics",
         "title": "舊聞", "url": "https://example.test/old",
         "published_at": "2026-01-01T00:00:00", "fetched_at": "2026-01-01T00:00:00"},
    ])
    before_bytes = path.read_bytes()
    before_cols = {row[1] for row in sqlite3_connect_columns(path)}

    tmp_path_copy = _safe_copy_for_dry_run(path)
    try:
        db = Database(tmp_path_copy)
        db.connect()
        try:
            result = backfill_topics(
                db,
                sources={},
                election_config=load_election_config(),
                election_entities=load_entities(),
                batch_size=1,
            )
            assert result["scanned"] == 1
            cols = {row[1] for row in db.conn.execute("PRAGMA table_info(articles)")}
            assert "identity_key" in cols
            assert "published_at_precision" in cols
            assert "delivery_eligible" in cols
        finally:
            db.close()
    finally:
        import shutil
        shutil.rmtree(tmp_path_copy.parent, ignore_errors=True)

    assert path.read_bytes() == before_bytes
    assert {row[1] for row in sqlite3_connect_columns(path)} == before_cols


def sqlite3_connect_columns(path):
    import sqlite3
    conn = sqlite3.connect(str(path))
    try:
        return conn.execute("PRAGMA table_info(articles)").fetchall()
    finally:
        conn.close()


def test_diagnosis_read_only_on_old_schema(tmp_path):
    path = tmp_path / "diag-old.db"
    _create_old_articles_db(path, [
        {"source_id": "s", "source_name": "S", "category": "politics",
         "title": "舊聞", "url": "https://example.test/old",
         "published_at": "2026-01-01T00:00:00", "fetched_at": "2026-01-01T00:00:00"},
    ])
    before_cols = sqlite3_connect_columns(path)
    before_row_count = sqlite3_count(path)

    from app.diagnose import run_diagnosis
    db = Database(path, read_only=True)
    db.connect()
    try:
        run_diagnosis([], db, tmp_path / "diagnostics")
        assert db.count_articles() == 1
    finally:
        db.close()

    assert {row[1] for row in sqlite3_connect_columns(path)} == {
        row[1] for row in before_cols
    }
    assert sqlite3_count(path) == before_row_count


def sqlite3_count(path):
    import sqlite3
    conn = sqlite3.connect(str(path))
    try:
        return conn.execute("SELECT COUNT(*) FROM articles").fetchone()[0]
    finally:
        conn.close()


def test_identity_migration_creates_index_before_backfill_and_preserves_rows(tmp_path):
    import sqlite3
    path = tmp_path / "identity-old.db"
    conn = sqlite3.connect(path)
    conn.execute("""
        CREATE TABLE articles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_id TEXT NOT NULL,
            source_name TEXT NOT NULL,
            category TEXT NOT NULL,
            title TEXT NOT NULL,
            url TEXT NOT NULL UNIQUE,
            published_at TEXT,
            fetched_at TEXT NOT NULL,
            position INTEGER NOT NULL
        )
    """)
    conn.execute(
        "INSERT INTO articles (source_id, source_name, category, title, url, "
        "published_at, fetched_at, position) VALUES (?,?,?,?,?,?,?,?)",
        ("udn", "聯合", "politics", "甲", "https://udn.com/news/story/6656/9635000",
         "2026-01-01T00:00:00", "2026-01-01T00:00:00", 1),
    )
    conn.execute(
        "INSERT INTO articles (source_id, source_name, category, title, url, "
        "published_at, fetched_at, position) VALUES (?,?,?,?,?,?,?,?)",
        ("udn", "聯合", "politics", "乙", "https://udn.com/news/story/7238/9635000",
         "2026-01-01T00:00:00", "2026-01-01T00:00:00", 2),
    )
    conn.commit()
    conn.close()

    db = Database(path)
    db.connect()
    try:
        cols = {row[1] for row in db.conn.execute("PRAGMA table_info(articles)")}
        assert "identity_key" in cols
        keys = [row[0] for row in db.conn.execute(
            "SELECT identity_key FROM articles ORDER BY id"
        )]
        assert keys[0] == keys[1]
        indexes = {
            row[1] for row in db.conn.execute("PRAGMA index_list(articles)")
        }
        assert "idx_articles_identity_key" in indexes
        assert db.count_articles() == 2
        urls = [row[0] for row in db.conn.execute("SELECT url FROM articles ORDER BY id")]
        assert urls == [
            "https://udn.com/news/story/6656/9635000",
            "https://udn.com/news/story/7238/9635000",
        ]
        collisions = db.conn.execute(
            "SELECT COUNT(*) FROM identity_collisions"
        ).fetchone()[0]
        assert collisions >= 1
    finally:
        db.close()

    # second migration is stable
    db2 = Database(path)
    db2.connect()
    try:
        assert db2.count_articles() == 2
        assert db2.schema_version() == SCHEMA_VERSION
    finally:
        db2.close()


def _fake_httpx_client(pages: dict[str, str]):
    import httpx

    class Client:
        def get(self, url: str):
            return httpx.Response(
                200,
                request=httpx.Request("GET", url),
                text=pages[url],
            )

        def close(self):
            pass

    return Client()


def test_nownews_collector_merges_banner_and_list_exact(monkeypatch):
    url = "https://www.nownews.com/cat/news-summary/military/"
    article_url = "https://www.nownews.com/news/123456"
    html = f"""
    <html><body>
      <a href="{article_url}" data-sec="banner_news">
        <h2 class="title">Banner no time</h2>
      </a>
      <ul id="ulNewsList">
        <li class="item">
          <a href="{article_url}">
            <h3 class="title">List exact</h3>
            <time datetime="2026-09-09 18:30">2026-09-09 18:30</time>
          </a>
        </li>
      </ul>
    </body></html>
    """
    cfg = {
        "id": "nownews_military",
        "name": "NOWnews",
        "type": "nownews_military",
        "category": "military",
        "topic": "military",
        "military_source_type": "commercial_military",
        "url": url,
        "max_pages": 1,
    }
    collector = NownewsMilitaryCollector(cfg)
    collector._client = _fake_httpx_client({url: html})
    monkeypatch.setattr(
        "app.collectors.military._now_taipei",
        lambda: datetime(2026, 9, 9, 19, 0, tzinfo=TAIPEI),
    )
    rows = collector.collect()
    assert len(rows) == 1
    assert rows[0].published_at_precision == "exact"
    assert rows[0].published_at == datetime(2026, 9, 9, 18, 30, tzinfo=TAIPEI)


def test_nownews_collector_merges_across_pages_without_first_win(monkeypatch):
    """page 1 exact must not be downgraded by page 2 banner unknown."""
    url = "https://www.nownews.com/cat/news-summary/military/"
    page1 = url.rstrip("/") + "/page/1/"
    article_url = "https://www.nownews.com/news/654321"
    exact_html = f"""
    <html><body><ul id="ulNewsList">
      <li class="item"><a href="{article_url}">
        <h3 class="title">Page one exact</h3>
        <time datetime="2026-09-09 18:30">2026-09-09 18:30</time>
      </a></li>
    </ul></body></html>
    """
    banner_html = f"""
    <html><body>
      <a href="{article_url}" data-sec="banner_news">
        <h2 class="title">Page two banner unknown</h2>
      </a>
      <ul id="ulNewsList"></ul>
    </body></html>
    """
    cfg = {
        "id": "nownews_military",
        "name": "NOWnews",
        "type": "nownews_military",
        "category": "military",
        "topic": "military",
        "military_source_type": "commercial_military",
        "url": url,
        "max_pages": 2,
    }
    collector = NownewsMilitaryCollector(cfg)
    collector._client = _fake_httpx_client({url: exact_html, page1: banner_html})
    monkeypatch.setattr(
        "app.collectors.military._now_taipei",
        lambda: datetime(2026, 9, 9, 19, 0, tzinfo=TAIPEI),
    )
    rows = collector.collect()
    assert len(rows) == 1
    assert rows[0].published_at_precision == "exact"
    assert rows[0].published_at == datetime(2026, 9, 9, 18, 30, tzinfo=TAIPEI)


def test_word_publish_time_respects_article_precision(tmp_path):
    run = datetime(2026, 9, 9, 19, 0, tzinfo=TAIPEI)
    date_only = article(
        "https://mna.test/date-only",
        "MNA date only",
        source_id="mna_military",
        category="military",
        published_at=datetime(2026, 9, 9, 0, 0, tzinfo=TAIPEI),
        precision="date_only",
    )
    exact = article(
        "https://media.test/exact",
        "Exact media",
        category="politics",
        published_at=datetime(2026, 9, 9, 18, 30, tzinfo=TAIPEI),
        precision="exact",
    )
    unknown = article(
        "https://media.test/unknown",
        "Unknown media",
        category="politics",
        published_at=None,
        precision="unknown",
    )
    output = build_word_digest(
        [date_only, exact, unknown],
        tmp_path,
        generated_at=run,
        military_topic_urls={date_only.url},
    )
    text = "\n".join(p.text for p in Document(output).paragraphs)
    assert "发布日期：2026-09-09" in text
    assert "发布时间：2026-09-09 18:30" in text
    assert "2026-09-09 00:00" not in text


def test_recovery_baseline_date_only_but_exact_unaffected(tmp_path):
    now = datetime(2026, 9, 9, 14, 0, tzinfo=TAIPEI)
    health = SourceHealthStore(tmp_path / "health.json")
    health.update(
        "mna_cont",
        SourceOutcome(200, True, 1),
        now=now - timedelta(minutes=30),
    )
    health.update(
        "mna_rec",
        SourceOutcome(200, True, 1),
        now=now - timedelta(days=2),
    )
    continuity = _get_source_continuity(
        health,
        [{"id": "mna_cont"}, {"id": "mna_rec"}],
        now,
        3,
    )
    assert continuity == {"mna_cont": True, "mna_rec": False}

    today_cont = _make_dated("https://mna.test/continuous", 9, source_id="mna_cont")
    today_rec = _make_dated("https://mna.test/recovery", 9, source_id="mna_rec")
    exact_after_outage = article(
        "https://mna.test/exact",
        "exact after outage",
        source_id="mna_rec",
        category="military",
        published_at=now - timedelta(minutes=30),
        precision="exact",
    )

    # First enable (baseline=0) still suppresses date-only.
    first = _classify_delivery_articles(
        [today_cont], {"mna_cont": 0}, now, source_continuity=continuity
    )
    assert first["date_only_eligible"] == []

    # Continuous source delivers new date-only.
    continuous = _classify_delivery_articles(
        [today_cont],
        {"mna_cont": 5},
        now,
        source_continuity=continuity,
    )
    assert continuous["date_only_eligible"] == [today_cont]

    # Recovery run holds back date-only for the source, but not exact news.
    recovery = _classify_delivery_articles(
        [today_rec, exact_after_outage],
        {"mna_rec": 5},
        now,
        source_continuity=continuity,
    )
    assert recovery["date_only_eligible"] == []
    assert recovery["fresh_articles"] == [exact_after_outage]
    assert today_rec in recovery["stale_articles"]

    # Next normal run after a successful health update is continuous again.
    health.update("mna_rec", SourceOutcome(200, True, 1), now=now)
    next_continuity = _get_source_continuity(
        health, [{"id": "mna_rec"}], now + timedelta(minutes=30), 3
    )
    next_delivery = _classify_delivery_articles(
        [today_rec],
        {"mna_rec": 5},
        now + timedelta(minutes=30),
        source_continuity=next_continuity,
    )
    assert next_continuity["mna_rec"] is True
    assert next_delivery["date_only_eligible"] == [today_rec]


def test_get_articles_between_preserves_metadata_and_word_date_only(tmp_path):
    db = Database(tmp_path / "between.db")
    db.connect()
    try:
        date_only = article(
            "https://mna.test/between",
            "between date only",
            source_id="mna_military",
            category="military",
            published_at=datetime(2026, 9, 9, 0, 0, tzinfo=TAIPEI),
            precision="date_only",
        )
        date_only.filter_reason = None
        db.save_articles([date_only])
        rows = db.get_articles_between(
            datetime(2026, 9, 1, tzinfo=TAIPEI),
            datetime(2026, 9, 30, tzinfo=TAIPEI),
            eligible_only=True,
        )
        assert len(rows) == 1
        assert rows[0].published_at_precision == "date_only"
        assert rows[0].delivery_eligible is True
        word = build_word_digest(rows, tmp_path, generated_at=datetime(2026, 9, 9, 12, 0, tzinfo=TAIPEI))
        text = "\n".join(p.text for p in Document(word).paragraphs)
        assert "发布日期：2026-09-09" in text
        assert "00:00" not in text
    finally:
        db.close()


def test_diagnose_accepts_production_type_and_legacy_collector(tmp_path, monkeypatch):
    import app.diagnose as diagnose_mod

    class FakeDiagnoseCollector:
        seen = []

        def __init__(self, source):
            self.source = source
            self.last_outcome = None

        def collect(self):
            self.seen.append(self.source.get("id"))
            return [
                article(
                    f"https://diag.test/{self.source.get('id')}",
                    "diagnosis article",
                    category="politics",
                    published_at=datetime(2026, 9, 9, 12, 0, tzinfo=TAIPEI),
                )
            ]

        def close(self):
            pass

    monkeypatch.setitem(diagnose_mod.COLLECTOR_MAP, "president_json", FakeDiagnoseCollector)
    monkeypatch.setitem(diagnose_mod.COLLECTOR_MAP, "rss", FakeDiagnoseCollector)
    out = tmp_path / "diagnostics"
    diagnose_mod.run_diagnosis(
        [
            {"id": "president_press", "name": "总统府", "type": "president_json"},
            {"id": "legacy", "name": "Legacy", "collector": "rss"},
            {"id": "unsupported", "name": "Unsupported", "type": "does_not_exist"},
        ],
        None,
        out,
        run_started_at=datetime(2026, 9, 9, 12, 0, tzinfo=TAIPEI),
    )
    assert set(FakeDiagnoseCollector.seen) == {"president_press", "legacy"}
    assert (out / "latest_collection.csv").exists()
    assert (out / "latest_diagnosis.md").exists()


def test_fix3_end_to_end_delivery_export_backfill_semantics(tmp_path):
    run = datetime(2026, 9, 9, 19, 0, tzinfo=TAIPEI)
    db = Database(tmp_path / "integration.db")
    db.connect()
    try:
        a = article(
            "https://example.test/a",
            "A normal politics",
            category="politics",
            published_at=run - timedelta(minutes=30),
            precision="exact",
        )
        b = article(
            "https://example.test/b",
            "B lottery",
            category="economy",
            published_at=run - timedelta(minutes=30),
            precision="exact",
        )
        b.delivery_eligible = False
        b.filter_reason = "content_filter"
        c = _make_dated(
            "https://mna.test/c", 9, source_id="mna_cont", title="C continuous date-only"
        )
        d = _make_dated(
            "https://mna.test/d", 9, source_id="mna_rec", title="D recovery date-only"
        )
        e_unknown = article(
            "https://www.nownews.com/news/999",
            "E banner",
            category="military",
            published_at=None,
            precision="unknown",
        )
        e_exact = article(
            "https://www.nownews.com/news/999",
            "E list exact",
            category="military",
            published_at=run - timedelta(minutes=30),
            precision="exact",
        )
        e_pair, _ = deduplicate_articles_by_url([e_unknown, e_exact])
        assert e_pair[0].published_at_precision == "exact"
        e = e_pair[0]

        assert db.save_articles([a, b, c, d, e]) == [a, b, c, d, e]

        delivery = run_delivery_core(
            [a, b, c, d, e],
            {"mna_cont": 5, "mna_rec": 5},
            run,
            international_config=None,
            importance_rules_config={"enabled": False, "thresholds": {}, "rules": []},
            prepare_international_delivery=lambda articles, _cfg: (
                list(articles),
                {article.url: [article] for article in articles},
            ),
            enrich_summaries=lambda _articles: None,
            excluded_delivery_urls={b.url},
            source_continuity={"mna_cont": True, "mna_rec": False},
            catch_up_enabled=False,
        )
        delivery_urls = {article.url for article in delivery.delivery_articles}
        assert a.url in delivery_urls
        assert b.url not in delivery_urls
        assert c.url in delivery_urls
        assert d.url not in delivery_urls
        assert e.url in delivery_urls

        word = build_word_digest(
            delivery.delivery_articles,
            tmp_path,
            generated_at=run,
            military_topic_urls={c.url, e.url},
        )
        word_text = "\n".join(p.text for p in Document(word).paragraphs)
        assert "发布日期：2026-09-09" in word_text
        assert "发布时间：2026-09-09 18:30" in word_text
        assert "2026-09-09 00:00" not in word_text

        export_articles = db.get_articles_since(
            datetime(2000, 1, 1), eligible_only=True
        )
        assert b.url not in {article.url for article in export_articles}

        between = db.get_articles_between(
            datetime(2026, 9, 1, tzinfo=TAIPEI),
            datetime(2026, 9, 30, tzinfo=TAIPEI),
            eligible_only=True,
        )
        assert b.url not in {article.url for article in between}
        c_loaded = next(article for article in between if article.url == c.url)
        assert c_loaded.published_at_precision == "date_only"
        backfill_word = build_word_digest(
            between, tmp_path, generated_at=run,
        )
        backfill_text = "\n".join(p.text for p in Document(backfill_word).paragraphs)
        assert "发布日期：2026-09-09" in backfill_text
        assert "2026-09-09 00:00" not in backfill_text
    finally:
        db.close()
