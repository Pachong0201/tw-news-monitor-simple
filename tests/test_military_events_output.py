"""Real SQLite/CLI/Word integration; no collector, HTTP or LLM calls."""
from copy import deepcopy
from datetime import datetime, timezone
import json
from pathlib import Path
import sqlite3
import subprocess
import sys

from docx import Document
import pytest

from app.database import Database
from app.importance import ImportanceResult
from app.military.cli import read_articles, main
from app.military.config import load_rules
from app.military.events import build_events
from app.military.output import markdown
from app.models import Article
from app.word_digest import build_word_digest

NOW = datetime(2026, 9, 5, 9)


def rows():
    return [Article("ltn_defense", "自由时报·军武", "politics", title, f"https://example.com/{i}", NOW, NOW, i)
            for i, title in enumerate(["F-16V进行新型飞弹挂载测试", "空军测试F16V新型导弹",
                                       "汉光演习今日展开", "海鲲号潜舰进行潜航测试", "台股收盘上涨"])]


@pytest.fixture
def db_path(tmp_path):
    path = tmp_path / "news.db"
    db = Database(path)
    db.connect()
    db.create_tables()
    db.save_articles(rows())
    db.close()
    return path


def test_reads_real_schema_preserves_ids_no_writes(db_path):
    before = db_path.read_bytes()
    articles, ids = read_articles(db_path, 24, NOW)
    assert len(articles) == 5
    assert set(ids.values()) == set(range(1, 6))
    assert db_path.read_bytes() == before


def test_timezones_future_and_missing_publication(db_path):
    with sqlite3.connect(db_path) as c:
        c.execute("UPDATE articles SET published_at='2026-09-05T01:00:00+00:00' WHERE id=1")
        c.execute("UPDATE articles SET published_at='2020-01-01T00:00:00' WHERE id=2")
        c.execute("UPDATE articles SET published_at='2026-09-06T00:00:00' WHERE id=3")
        c.execute("UPDATE articles SET published_at=NULL WHERE id=4")
    data, ids = read_articles(db_path, 24, datetime(2026, 9, 5, 1, tzinfo=timezone.utc))
    assert set(ids.values()) == {1, 4, 5}


def test_invalid_stored_time_skipped(db_path, caplog):
    with sqlite3.connect(db_path) as c:
        c.execute("UPDATE articles SET published_at='not-a-date' WHERE id=1")
    data, _ = read_articles(db_path, 24, NOW)
    assert len(data) == 4
    assert "invalid timestamp" in caplog.text


def test_cli_json_and_markdown(db_path, tmp_path):
    output = tmp_path / "events.json"
    assert main(["--db", str(db_path), "--now", NOW.isoformat(), "--output", str(output)]) == 0
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["input_articles"] == 5
    assert len(payload["events"]) == 3
    assert payload["events"][0]["article_ids"]
    result = tmp_path / "events.md"
    assert main(["--db", str(db_path), "--now", NOW.isoformat(), "--format", "markdown", "--output", str(result)]) == 0
    assert "# 军武动态" in result.read_text(encoding="utf-8")


def test_cli_subprocess_real_entrypoint(db_path):
    result = subprocess.run([sys.executable, "-m", "app.military.cli", "--db", str(db_path), "--now", NOW.isoformat()], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr
    assert len(json.loads(result.stdout)["events"]) == 3


def test_cli_missing_database_does_not_create(tmp_path):
    path = tmp_path / "absent.db"
    assert main(["--db", str(path)]) == 2
    assert not path.exists()


@pytest.mark.parametrize("value", ["0", "-1", "nan", "inf"])
def test_cli_invalid_hours(db_path, value):
    assert main(["--db", str(db_path), "--hours", value]) == 2


def test_cli_cannot_overwrite_database_or_rules(db_path):
    before = db_path.read_bytes()
    assert main(["--db", str(db_path), "--output", str(db_path)]) == 2
    assert db_path.read_bytes() == before


def test_markdown_complete_and_empty():
    text = markdown(build_events(rows(), load_rules()))
    for value in ["类别：", "重要度：", "来源：", "核心实体：", "代表报道：", "研判依据：", "相关新闻：2 篇"]:
        assert value in text
    assert "本时间段没有识别到" in markdown([])


def test_word_events_real_routing_and_no_debug(tmp_path):
    rules = load_rules()
    rules["word_enabled"] = True
    data = rows()
    out = build_word_digest(data, tmp_path, generated_at=NOW, military_event_config=rules)
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["一、军武动态", "二、政治新闻"]
    assert text.count("相关新闻：2 篇") == 1
    assert "台股收盘上涨" in text
    assert "为何重要：" in text and "台军装备" in text
    assert "cluster_score" not in text and "title_similarity" not in text
    links = [r.target_ref for r in doc.part.rels.values() if "hyperlink" in r.reltype]
    assert len(links) >= 3


def test_word_disabled_identical_legacy(tmp_path):
    data = rows()
    a = build_word_digest(data, tmp_path / "a", generated_at=NOW)
    b = build_word_digest(data, tmp_path / "b", generated_at=NOW, military_event_config=load_rules())
    assert a.read_bytes() == b.read_bytes()


def test_word_failure_preserves_existing_topic(tmp_path, caplog):
    out = build_word_digest(rows(), tmp_path, military_topic_urls={rows()[0].url}, military_event_config={"word_enabled": True})
    assert "军武动态" in "\n".join(p.text for p in Document(out).paragraphs)
    assert "Military event" in caplog.text


def test_word_highlights_unmodified_and_unknown_legacy_kept(tmp_path):
    data = rows()
    rules = load_rules()
    rules["word_enabled"] = True
    unknown = Article("ydn_defense_focus", "青年日报·国防焦点", "politics", "部队组织经验交流座谈", "https://example.com/unknown", NOW, NOW, 1)
    out = build_word_digest(data + [unknown], tmp_path, military_event_config=rules,
                            military_topic_urls={unknown.url},
                            importance_results=[(data[0], ImportanceResult(90, "critical", ["existing"]))])
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "重点提示" in text and "【重大】" in text
    assert unknown.title in text
    assert "相关新闻：2 篇" in text


def test_cli_word_real_generation(db_path, tmp_path):
    output = tmp_path / "word"
    assert main(["--db", str(db_path), "--now", NOW.isoformat(), "--word-dir", str(output), "--output", str(tmp_path / "events.json")]) == 0
    paths = list(output.glob("*.docx"))
    assert len(paths) == 1
    assert any("相关新闻：2 篇" in p.text for p in Document(paths[0]).paragraphs)
