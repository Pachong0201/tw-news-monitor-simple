"""军武专题栏仅收录军武来源（来源白名单）。

回归保护：普通政治/经济频道的军武标题稿、以及 DB 中陈旧/污染的
military topic 记录，都不得进入 Word "军武动态" 专题栏；被排除的
文章必须仍出现在其原有分类栏（不消失）。
"""
from datetime import datetime

from docx import Document
import pytest

from app.election2026.config import load_election_config, load_entities
from app.military import is_military_source_article, military_source_ids
from app.military.config import load_rules
from app.models import Article
from app.word_digest import build_word_digest

NOW = datetime(2026, 9, 5, 9, 0)
ELECTION_CONFIG = load_election_config()
ELECTION_ENTITIES = load_entities()


def article(title, url, source_id="ltn_defense", source_name="自由时报·军武",
            category="politics"):
    return Article(source_id=source_id, source_name=source_name, category=category,
                   title=title, url=url, published_at=NOW, fetched_at=NOW, position=1)


def render(tmp_path, articles, *, military_urls=None, military_event_config=None,
           election=False):
    output = build_word_digest(
        articles, tmp_path, generated_at=NOW,
        military_topic_urls=set(military_urls or []),
        military_event_config=military_event_config,
        election_config=ELECTION_CONFIG if election else None,
        election_entities=ELECTION_ENTITIES if election else None,
    )
    return Document(output)


def texts(doc):
    return [p.text for p in doc.paragraphs]


def headings(doc):
    return [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]


# ---- helper 单元测试 -------------------------------------------------------

def test_military_source_ids_cover_all_military_sources():
    ids = military_source_ids()
    assert {"ltn_defense", "udn_military", "nownews_military", "ydn_defense_focus",
            "ydn_weapons_tour", "ydn_military_world", "mna_military"} <= ids
    assert "cna_politics" not in ids
    assert "ltn_politics" not in ids


def test_is_military_source_article_accepts_only_military_sources():
    assert is_military_source_article(article("任何", "https://example.com/1"))
    assert is_military_source_article(
        article("任何", "https://example.com/2", source_id="mna_military",
                source_name="国防部军事新闻通讯社"))
    assert not is_military_source_article(
        article("任何", "https://example.com/3", source_id="cna_politics",
                source_name="中央社"))
    assert not is_military_source_article(
        article("任何", "https://example.com/4", source_id="ltn_politics",
                source_name="自由时报"))
    assert not is_military_source_article(
        Article(source_id=None, source_name="x", category="politics", title="t",
                url="https://example.com/5", published_at=NOW, fetched_at=NOW,
                position=1))


# ---- Word 专题栏：白名单拦截 ----------------------------------------------

def test_general_source_military_headline_never_enters_military_section(tmp_path):
    general = article("F-16V军购案立法院审议", "https://example.com/general",
                      source_id="cna_politics", source_name="中央社")
    doc = render(tmp_path, [general])
    flat = "\n".join(texts(doc))
    # 无军武来源文章 → 不渲染军武专题栏
    assert "一、军武动态" not in headings(doc)
    # 政治稿仍出现在政治栏（不消失）
    assert "F-16V军购案立法院审议" in flat
    assert "一、政治新闻" in headings(doc)


def test_military_source_article_enters_military_section(tmp_path):
    military = article("国军防空飞弹战备操演", "https://example.com/mil")
    doc = render(tmp_path, [military], military_urls={military.url})
    assert headings(doc) == ["一、军武动态"]
    assert "国军防空飞弹战备操演" in "\n".join(texts(doc))


def test_stale_topic_of_general_source_is_excluded_and_not_lost(tmp_path):
    # 模拟 DB 污染：general 来源 URL 也带 military topic（历史/补发记录）
    general = article("F-16V军购案立法院审议", "https://example.com/polluted",
                      source_id="cna_politics", source_name="中央社")
    doc = render(tmp_path, [general], military_urls={general.url})
    flat = "\n".join(texts(doc))
    # 不进军武专题栏
    assert "一、军武动态" not in headings(doc)
    # 仍出现在政治栏
    assert general.title in flat
    assert "一、政治新闻" in headings(doc)


def test_stale_topic_excluded_but_military_source_still_rendered(tmp_path):
    military = article("潜舰海试完成", "https://example.com/kept")
    general = article("F-16V军购案立法院审议", "https://example.com/dropped",
                      source_id="cna_politics", source_name="中央社")
    doc = render(tmp_path, [military, general],
                 military_urls={military.url, general.url})
    headings_out = headings(doc)
    flat = "\n".join(texts(doc))
    assert "一、军武动态" in headings_out
    assert military.title in flat
    assert general.title in flat
    assert any("政治新闻" in h for h in headings_out)
    # general 稿只出现一次（政治栏），未进军事栏
    assert flat.count(general.title) == 1


def test_events_route_ignores_general_source_even_when_classifiable(tmp_path):
    rules = load_rules()
    rules["word_enabled"] = True
    general = article("解放军军机绕台与实弹演习", "https://example.com/pla",
                      source_id="cna_politics", source_name="中央社")
    doc = render(tmp_path, [general], military_event_config=rules)
    flat = "\n".join(texts(doc))
    assert "一、军武动态" not in headings(doc)
    assert general.title in flat  # 回到政治栏
    assert "相关新闻" not in flat


def test_events_route_still_groups_military_source_articles(tmp_path):
    rules = load_rules()
    rules["word_enabled"] = True
    military = article("F-16V进行新型飞弹挂载测试", "https://example.com/a")
    other = article("空军测试F16V新型导弹", "https://example.com/b")
    doc = render(tmp_path, [military, other], military_event_config=rules)
    flat = "\n".join(texts(doc))
    assert headings(doc) == ["一、军武动态"]
    assert "相关新闻：2 篇" in flat


def test_mixed_sources_events_include_only_military_members(tmp_path):
    rules = load_rules()
    rules["word_enabled"] = True
    military = article("海鲲号潜舰进行潜航测试", "https://example.com/a")
    general = article("海鲲号潜舰交付国会听证会", "https://example.com/b",
                      source_id="cna_politics", source_name="中央社")
    doc = render(tmp_path, [military, general], military_event_config=rules)
    flat = "\n".join(texts(doc))
    assert "一、军武动态" in headings(doc)
    assert military.title in flat
    assert general.title in flat  # 普通栏可见
    assert any("政治新闻" in h for h in headings(doc))


# ---- fail-closed：白名单不可用 --------------------------------------------

def test_unreadable_sources_config_disables_military_section(tmp_path, monkeypatch):
    from app import military as military_module

    def broken():
        raise OSError("no sources.yaml")
    monkeypatch.setattr(military_module, "military_source_ids", broken)
    military = article("国军防空飞弹战备操演", "https://example.com/mil")
    doc = render(tmp_path, [military], military_urls={military.url})
    # fail-closed：不再渲染专题军武栏，也不抛错
    assert "一、军武动态" not in headings(doc)
