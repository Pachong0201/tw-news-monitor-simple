from datetime import datetime

from docx import Document

from app.election2026.config import load_election_config, load_entities
from app.importance import ImportanceResult
from app.models import Article
from app.word_digest import build_word_digest


NOW = datetime(2026, 9, 4, 9, 0)
ELECTION_CONFIG = load_election_config()
ELECTION_ENTITIES = load_entities()


def article(title, url, category="politics", source_id="udn", source_name="联合新闻网"):
    return Article(
        source_id=source_id,
        source_name=source_name,
        category=category,
        title=title,
        url=url,
        published_at=NOW,
        fetched_at=NOW,
        position=1,
    )


def render(tmp_path, articles, *, military_urls=None, importance=None, election=True):
    output = build_word_digest(
        articles,
        tmp_path,
        generated_at=NOW,
        military_topic_urls=set(military_urls or []),
        importance_results=importance,
        election_config=ELECTION_CONFIG if election else None,
        election_entities=ELECTION_ENTITIES if election else None,
    )
    doc = Document(output)
    return doc, [paragraph.text for paragraph in doc.paragraphs]


def level_one_headings(doc):
    return [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style.name == "Heading 1"
    ]


def test_no_military_articles_hides_section(tmp_path):
    doc, _ = render(
        tmp_path,
        [article("立法院审查总预算", "https://example.com/politics")],
        military_urls=set(),
    )
    assert not any("军武动态" in heading for heading in level_one_headings(doc))


def test_military_section_is_rendered(tmp_path):
    military = article("国军完成防空飞弹战备操演", "https://example.com/military")
    doc, texts = render(tmp_path, [military], military_urls={military.url})
    assert level_one_headings(doc) == ["一、军武动态"]
    assert sum(military.title in text for text in texts) == 1


def test_election_and_military_overlap_appears_in_both_sections(tmp_path):
    overlap = article(
        "台南市长参选人提出国防产业政策",
        "https://example.com/overlap",
    )
    doc, texts = render(tmp_path, [overlap], military_urls={overlap.url})
    assert level_one_headings(doc) == ["一、九合一选举", "二、军武动态"]
    assert sum(overlap.title in text for text in texts) == 2


def test_critical_military_appears_in_highlights_and_military(tmp_path):
    military = article("共军军演与飞弹动态", "https://example.com/critical")
    importance = [
        (military, ImportanceResult(score=90, level="critical", matched_rules=["x"]))
    ]
    doc, texts = render(
        tmp_path,
        [military],
        military_urls={military.url},
        importance=importance,
        election=False,
    )
    assert level_one_headings(doc) == ["一、重点提示", "二、军武动态"]
    assert sum(military.title in text for text in texts) == 2


def test_ordinary_military_is_not_duplicated_into_general_sections(tmp_path):
    military = article("潜舰完成海试", "https://example.com/submarine", "international")
    politics = article("立法院审查法案", "https://example.com/law")
    doc, texts = render(
        tmp_path,
        [military, politics],
        military_urls={military.url},
        election=False,
    )
    assert level_one_headings(doc) == ["一、军武动态", "二、政治新闻"]
    assert sum(military.title in text for text in texts) == 1


def test_all_top_level_sections_are_contiguous_in_target_order(tmp_path):
    official = article(
        "总统府发布新闻",
        "https://example.com/official",
        source_id="president_press",
        source_name="台湾总统府",
    )
    highlight = article("重大政策发布", "https://example.com/highlight")
    election = article("民进党启动2026县市长提名机制", "https://example.com/election")
    military = article("战机完成战备任务", "https://example.com/military")
    politics = article("立法院审查法案", "https://example.com/politics")
    economy = article("台股收盘上涨", "https://example.com/economy", "economy")
    international = article("两岸交流议题受关注", "https://example.com/intl", "international")
    importance = [
        (highlight, ImportanceResult(score=70, level="important", matched_rules=["x"]))
    ]
    doc, _ = render(
        tmp_path,
        [official, highlight, election, military, politics, economy, international],
        military_urls={military.url},
        importance=importance,
    )
    assert level_one_headings(doc) == [
        "一、官方信源",
        "二、重点提示",
        "三、九合一选举",
        "四、军武动态",
        "五、政治新闻",
        "六、经济新闻",
        "七、国际及两岸新闻",
    ]
